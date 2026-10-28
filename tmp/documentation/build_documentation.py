from __future__ import annotations

import argparse
import html
import json
import re
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Image as PDFImage
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / 'docs' / 'PROJECT_DOCUMENTATION_EL.md'
WORK = ROOT / 'tmp' / 'documentation'
DOCX = ROOT / 'output' / 'docs' / 'Property_Co_Pilot_Technical_Documentation_EL.docx'
PDF = ROOT / 'output' / 'pdf' / 'Property_Co_Pilot_Technical_Documentation_EL.pdf'
FONT_ROOT = Path('C:/Windows/Fonts')
INK = '#171717'
ACCENT = '#596044'
PAGES = SOURCE.read_text(encoding='utf-8').split('@page')


def blocks(text):
    lines = text.strip().splitlines()
    result = []
    index = 0
    while index < len(lines):
        line = lines[index].strip()
        if not line or line == '@cover':
            index += 1
            continue
        if line.startswith('```'):
            code = []
            index += 1
            while index < len(lines) and not lines[index].startswith('```'):
                code.append(lines[index])
                index += 1
            result.append(('code', code))
        elif line.startswith('|'):
            rows = []
            while index < len(lines) and lines[index].strip().startswith('|'):
                row = [cell.strip() for cell in lines[index].strip().strip('|').split('|')]
                if not all(re.fullmatch(r'[-: ]+', cell) for cell in row):
                    rows.append(row)
                index += 1
            result.append(('table', rows))
            continue
        elif line.startswith('@equation '):
            result.append(('equation', line[10:]))
        elif line.startswith('@figure '):
            result.append(('figure', line[8:]))
        elif line == '@toc':
            result.append(('toc', None))
        elif line.startswith('## '):
            result.append(('h2', line[3:]))
        elif line.startswith('# '):
            result.append(('h1', line[2:]))
        else:
            paragraph = [line]
            while index + 1 < len(lines) and lines[index + 1].strip():
                index += 1
                paragraph.append(lines[index].strip())
            result.append(('p', ' '.join(paragraph)))
        index += 1
    return result


PARSED = [blocks(page) for page in PAGES]
TOC = []
for page_index, page in enumerate(PARSED, 1):
    for kind, value in page:
        if kind == 'h1' and page_index not in (1, 3):
            TOC.append((value, page_index, f'section_{page_index}'))


def font(name, size):
    return ImageFont.truetype(str(FONT_ROOT / name), size)


def centered(draw, bounds, text, size=30):
    x, y, width, height = bounds
    selected = font('calibri.ttf', size)
    box = draw.multiline_textbbox((0, 0), text, font=selected, spacing=8, align='center')
    draw.multiline_text((x + (width - box[2]) / 2, y + (height - box[3]) / 2 - box[1]), text, font=selected, fill=INK, spacing=8, align='center')


def arrow(draw, start, end):
    import math

    draw.line((start, end), fill=ACCENT, width=4)
    angle = math.atan2(end[1] - start[1], end[0] - start[0])
    sides = [
        (end[0] - 16 * math.cos(angle - 0.45), end[1] - 16 * math.sin(angle - 0.45)),
        (end[0] - 16 * math.cos(angle + 0.45), end[1] - 16 * math.sin(angle + 0.45)),
    ]
    draw.polygon([end, *sides], fill=ACCENT)


def make_figures():
    targets = {}
    image = Image.new('RGB', (1320, 590), 'white')
    draw = ImageDraw.Draw(image)
    boxes = [
        (15, 35, 320, 110, 'React / TypeScript\nBrowser'),
        (500, 35, 320, 110, 'Go API\nONNX inference'),
        (985, 35, 320, 110, 'PostgreSQL\nPostGIS'),
        (15, 390, 320, 110, 'Tree.py\nTraining process'),
        (500, 390, 320, 110, 'Model artifacts\nONNX · UBJ · manifest'),
        (985, 390, 320, 110, 'FastAPI\nNative explanations'),
    ]
    for x, y, width, height, label in boxes:
        draw.rounded_rectangle((x, y, x + width, y + height), radius=10, fill='#F4F4EF', outline='#909482', width=2)
        centered(draw, (x, y, width, height), label, 31)
    for start, end in [((335, 90), (500, 90)), ((820, 90), (985, 90)), ((535, 145), (175, 390)), ((785, 145), (1145, 390)), ((335, 445), (500, 445)), ((820, 445), (985, 445)), ((660, 390), (660, 145))]:
        arrow(draw, start, end)
    for bounds, label in [((350, 38, 140, 35), 'HTTP'), ((835, 38, 135, 35), 'SQL'), ((190, 220, 290, 50), 'subprocess'), ((885, 220, 255, 50), 'HTTP'), ((675, 242, 155, 40), 'ONNX'), ((836, 470, 140, 40), 'UBJ')]:
        centered(draw, bounds, label, 26)
    targets['architecture'] = (image, 'Σχήμα 2.1  Υπηρεσίες και διαδρομές δεδομένων', 7.2)

    for key, labels, caption in [
        ('training', ['CSV\n29 features', 'Attr model\nresiduals', 'GeoPremium\n50 / 300 m', 'Phase2\n30 features', 'Artifacts\nSHAP / DB'], 'Σχήμα 3.1  Η αλυσίδα καθολικής εκπαίδευσης'),
        ('reload', ['Tree.py\nUBJ + hash', 'Go\nέλεγχος hash', 'Python\nload + έλεγχος', 'Python\nregistry swap', 'Go\nONNX swap'], 'Σχήμα 3.2  Η ακολουθία ελεγχόμενου reload'),
    ]:
        image = Image.new('RGB', (1500, 210), 'white')
        draw = ImageDraw.Draw(image)
        for index, label in enumerate(labels):
            x = 10 + index * 302
            draw.rounded_rectangle((x, 35, x + 265, 175), radius=10, fill='#F4F4EF', outline='#909482', width=2)
            centered(draw, (x, 35, 265, 140), label, 30)
            if index < len(labels) - 1:
                arrow(draw, (x + 266, 105), (x + 300, 105))
        targets[key] = (image, caption, 2.35)
    figures = {}
    for key, (image, caption, height) in targets.items():
        path = WORK / f'{key}.png'
        image.save(path)
        figures[key] = (path, caption, height)
    return figures


FIGURES = make_figures()


def bookmark(paragraph, name, number):
    start = OxmlElement('w:bookmarkStart')
    start.set(qn('w:id'), str(number))
    start.set(qn('w:name'), name)
    end = OxmlElement('w:bookmarkEnd')
    end.set(qn('w:id'), str(number))
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def internal_link(paragraph, label, anchor):
    element = OxmlElement('w:hyperlink')
    element.set(qn('w:anchor'), anchor)
    run = OxmlElement('w:r')
    text = OxmlElement('w:t')
    text.text = label
    run.append(text)
    element.append(run)
    paragraph._p.append(element)


def set_cell_properties(cell, header=False):
    properties = cell._tc.get_or_add_tcPr()
    margins = OxmlElement('w:tcMar')
    for side in ('top', 'left', 'bottom', 'right'):
        margin = OxmlElement(f'w:{side}')
        margin.set(qn('w:w'), '80')
        margin.set(qn('w:type'), 'dxa')
        margins.append(margin)
    properties.append(margins)
    if header:
        shading = OxmlElement('w:shd')
        shading.set(qn('w:fill'), 'E9EBE3')
        properties.append(shading)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def make_docx():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.05)
    section.right_margin = Cm(2.05)
    section.header_distance = Cm(0.85)
    section.footer_distance = Cm(0.8)
    section.different_first_page_header_footer = True
    styles = doc.styles
    for name in ['Normal', 'Body Text']:
        styles[name].font.name = 'Georgia'
        styles[name].font.size = Pt(11)
        styles[name].font.color.rgb = RGBColor.from_string('171717')
        styles[name].paragraph_format.line_spacing = 1.23
        styles[name].paragraph_format.space_after = Pt(9)
        styles[name].paragraph_format.widow_control = True
    for name, size in [('Title', 32), ('Subtitle', 19), ('Heading 1', 20), ('Heading 2', 14), ('Caption', 9)]:
        styles[name].font.name = 'Calibri'
        styles[name].font.size = Pt(size)
        styles[name].font.color.rgb = RGBColor.from_string('000000')
        styles[name].paragraph_format.space_before = Pt(12 if name == 'Heading 2' else 0)
        styles[name].paragraph_format.space_after = Pt(11)
        styles[name].paragraph_format.keep_with_next = name.startswith('Heading')
    styles['Caption'].font.color.rgb = RGBColor.from_string('596044')
    header = section.header.paragraphs[0]
    header.text = 'Property Co Pilot  |  Τεχνική τεκμηρίωση'
    header.runs[0].font.name = 'Calibri'
    header.runs[0].font.size = Pt(9)
    header.runs[0].font.color.rgb = RGBColor.from_string('666666')
    footer = section.footer.paragraphs[0]
    footer.paragraph_format.tab_stops.add_tab_stop(Cm(16.9), WD_TAB_ALIGNMENT.RIGHT)
    footer.add_run('Οκτώβριος 2026\t')
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    for run in footer.runs:
        run.font.name = 'Calibri'
        run.font.size = Pt(9)
    doc.core_properties.title = 'Property Co Pilot — Τεχνική τεκμηρίωση'
    doc.core_properties.subject = 'Αρχιτεκτονική και εξηγήσιμη ανάλυση ακινήτων'
    doc.core_properties.author = ''
    doc.core_properties.keywords = 'Property Co Pilot, XGBoost, SHAP, Go, PostGIS'
    language = OxmlElement('w:lang')
    language.set(qn('w:val'), 'el-GR')
    styles['Normal'].element.get_or_add_rPr().append(language)

    for page_number, page in enumerate(PARSED, 1):
        if page_number > 1:
            doc.add_page_break()
        else:
            blank = doc.add_paragraph()
            blank.paragraph_format.space_after = Cm(4.0)
        for kind, value in page:
            if kind in ('h1', 'h2'):
                style = 'Title' if page_number == 1 and kind == 'h1' else 'Subtitle' if page_number == 1 else 'Heading 1' if kind == 'h1' else 'Heading 2'
                paragraph = doc.add_paragraph(value, style)
                if kind == 'h1':
                    bookmark(paragraph, f'section_{page_number}', page_number)
                if page_number == 1:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    paragraph.paragraph_format.space_after = Pt(25)
            elif kind == 'p':
                paragraph = doc.add_paragraph(value)
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if page_number == 1 else WD_ALIGN_PARAGRAPH.JUSTIFY
                if page_number == 1:
                    paragraph.paragraph_format.space_after = Pt(22)
            elif kind == 'toc':
                for label, target_page, anchor in TOC:
                    paragraph = doc.add_paragraph()
                    paragraph.paragraph_format.space_after = Pt(11)
                    internal_link(paragraph, label, anchor)
                doc.add_paragraph('Οι τίτλοι λειτουργούν ως σύνδεσμοι προς τα κεφάλαια. Η πλοήγηση του Word διατίθεται επίσης από το παράθυρο επικεφαλίδων.')
            elif kind == 'table':
                table = doc.add_table(rows=len(value), cols=len(value[0]))
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                table.autofit = False
                widths = [6.2, 10.7] if len(value[0]) == 2 else [6.0, 5.6, 5.3] if len(value[0]) == 3 else [4.6, 3.8, 4.2, 4.3]
                for column, width in zip(table.columns, widths):
                    column.width = Cm(width)
                props = table._tbl.tblPr
                borders = OxmlElement('w:tblBorders')
                for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
                    border = OxmlElement(f'w:{side}')
                    border.set(qn('w:val'), 'single')
                    border.set(qn('w:sz'), '4')
                    border.set(qn('w:color'), 'C9CDC3')
                    borders.append(border)
                props.append(borders)
                header_row = OxmlElement('w:tblHeader')
                table.rows[0]._tr.get_or_add_trPr().append(header_row)
                for row_number, row in enumerate(value):
                    no_split = OxmlElement('w:cantSplit')
                    table.rows[row_number]._tr.get_or_add_trPr().append(no_split)
                    for cell, cell_text, width in zip(table.rows[row_number].cells, row, widths):
                        cell.width = Cm(width)
                        set_cell_properties(cell, row_number == 0)
                        paragraph = cell.paragraphs[0]
                        paragraph.paragraph_format.space_after = Pt(1)
                        paragraph.paragraph_format.line_spacing = 1.07
                        run = paragraph.add_run(cell_text)
                        run.font.name = 'Calibri'
                        run.font.size = Pt(10)
                        run.bold = row_number == 0
                after = doc.add_paragraph()
                after.paragraph_format.space_after = Pt(2)
                after.paragraph_format.line_spacing = 0.4
            elif kind == 'code':
                for line in value:
                    paragraph = doc.add_paragraph()
                    paragraph.paragraph_format.space_after = Pt(3)
                    paragraph.paragraph_format.line_spacing = 1.0
                    run = paragraph.add_run(line)
                    run.font.name = 'Consolas'
                    run.font.size = Pt(8.5)
            elif kind == 'equation':
                paragraph = doc.add_paragraph()
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                equation = OxmlElement('m:oMath')
                run = OxmlElement('m:r')
                text = OxmlElement('m:t')
                text.text = value
                run.append(text)
                equation.append(run)
                paragraph._p.append(equation)
            elif kind == 'figure':
                path, caption, height = FIGURES[value]
                paragraph = doc.add_paragraph()
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                picture = paragraph.add_run().add_picture(str(path), width=Cm(16.8))
                picture._inline.docPr.set('descr', caption)
                doc.add_paragraph(caption, 'Caption')
    DOCX.parent.mkdir(parents=True, exist_ok=True)
    doc.save(DOCX)
    print(f'DOCX {DOCX}')


def register_fonts():
    for name, filename in [('Georgia', 'georgia.ttf'), ('GeorgiaBold', 'georgiab.ttf'), ('GeorgiaItalic', 'georgiai.ttf'), ('Calibri', 'calibri.ttf'), ('CalibriBold', 'calibrib.ttf'), ('Consolas', 'consola.ttf')]:
        pdfmetrics.registerFont(TTFont(name, str(FONT_ROOT / filename)))
    pdfmetrics.registerFontFamily('Georgia', normal='Georgia', bold='GeorgiaBold', italic='GeorgiaItalic', boldItalic='GeorgiaBold')
    pdfmetrics.registerFontFamily('Calibri', normal='Calibri', bold='CalibriBold', italic='Calibri', boldItalic='CalibriBold')


def make_pdf():
    register_fonts()
    page_width, page_height = A4
    margin = 58
    content_width = page_width - 2 * margin
    styles = {
        'p': ParagraphStyle('body', fontName='Georgia', fontSize=11, leading=15.4, textColor=colors.HexColor(INK), alignment=TA_JUSTIFY, spaceAfter=10),
        'h1': ParagraphStyle('heading1', fontName='CalibriBold', fontSize=20, leading=24, spaceAfter=17, keepWithNext=True),
        'h2': ParagraphStyle('heading2', fontName='CalibriBold', fontSize=14, leading=18, spaceBefore=7, spaceAfter=11, keepWithNext=True),
        'title': ParagraphStyle('title', fontName='CalibriBold', fontSize=34, leading=40, alignment=TA_CENTER, spaceAfter=24),
        'subtitle': ParagraphStyle('subtitle', fontName='Calibri', fontSize=21, leading=29, alignment=TA_CENTER, spaceAfter=35),
        'cover': ParagraphStyle('cover', fontName='Georgia', fontSize=12, leading=20, alignment=TA_CENTER, spaceAfter=22),
        'cell': ParagraphStyle('cell', fontName='Calibri', fontSize=10.2, leading=12.4),
        'headercell': ParagraphStyle('headercell', fontName='CalibriBold', fontSize=10.2, leading=12.4),
        'caption': ParagraphStyle('caption', fontName='Calibri', fontSize=9, leading=12, textColor=colors.HexColor(ACCENT), spaceAfter=13),
        'equation': ParagraphStyle('equation', fontName='Georgia', fontSize=13, leading=21, alignment=TA_CENTER, spaceBefore=3, spaceAfter=13),
        'code': ParagraphStyle('code', fontName='Consolas', fontSize=8.4, leading=11.5, spaceAfter=2),
        'toc': ParagraphStyle('toc', fontName='Calibri', fontSize=12, leading=18, spaceAfter=11),
    }
    story = []
    expected_pages = {}
    for page_number, page in enumerate(PARSED, 1):
        if page_number > 1:
            story.append(PageBreak())
        else:
            story.append(Spacer(1, 123))
        first_heading = True
        for kind, value in page:
            if kind in ('p', 'h1', 'h2'):
                style = styles['title'] if page_number == 1 and kind == 'h1' else styles['subtitle'] if page_number == 1 and kind == 'h2' else styles['cover'] if page_number == 1 else styles[kind]
                text = html.escape(value)
                if kind in ('h1', 'h2') and first_heading:
                    text = f'<a name="section_{page_number}"/>' + text
                    first_heading = False
                    expected_pages[page_number] = value
                story.append(Paragraph(text, style))
            elif kind == 'toc':
                for label, target_page, anchor in TOC:
                    row = Table([[Paragraph(f'<link href="#{anchor}" color="{INK}">{html.escape(label)}</link>', styles['toc']), Paragraph(str(target_page), styles['toc'])]], colWidths=[content_width - 35, 35])
                    row.setStyle(TableStyle([('VALIGN', (0, 0), (-1, -1), 'TOP'), ('LEFTPADDING', (0, 0), (-1, -1), 0), ('RIGHTPADDING', (0, 0), (-1, -1), 0)]))
                    story.append(row)
            elif kind == 'table':
                cols = len(value[0])
                widths = [0.365, 0.635] if cols == 2 else [0.355, 0.332, 0.313] if cols == 3 else [0.275, 0.225, 0.25, 0.25]
                rows = [[Paragraph(html.escape(cell), styles['headercell'] if row_index == 0 else styles['cell']) for cell in row] for row_index, row in enumerate(value)]
                table = Table(rows, colWidths=[content_width * width for width in widths], repeatRows=1, hAlign='LEFT')
                table.setStyle(TableStyle([('GRID', (0, 0), (-1, -1), 0.4, colors.HexColor('#C9CDC3')), ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#E9EBE3')), ('VALIGN', (0, 0), (-1, -1), 'TOP'), ('LEFTPADDING', (0, 0), (-1, -1), 7), ('RIGHTPADDING', (0, 0), (-1, -1), 7), ('TOPPADDING', (0, 0), (-1, -1), 5), ('BOTTOMPADDING', (0, 0), (-1, -1), 5)]))
                story.append(table)
                story.append(Spacer(1, 12))
            elif kind == 'equation':
                story.append(Paragraph(html.escape(value), styles['equation']))
            elif kind == 'code':
                for line in value:
                    story.append(Paragraph(html.escape(line), styles['code']))
                story.append(Spacer(1, 7))
            elif kind == 'figure':
                path, caption, height = FIGURES[value]
                with Image.open(path) as picture:
                    aspect = picture.height / picture.width
                story.append(PDFImage(str(path), width=content_width, height=content_width * aspect))
                story.append(Paragraph(caption, styles['caption']))

    def decorate(canvas, document):
        canvas.setTitle('Property Co Pilot — Τεχνική τεκμηρίωση')
        canvas.setAuthor('')
        if document.page > 1:
            canvas.setFont('Calibri', 9)
            canvas.setFillColor(colors.HexColor('#666666'))
            canvas.drawString(margin, page_height - 31, 'Property Co Pilot  |  Τεχνική τεκμηρίωση')
            canvas.drawString(margin, 25, 'Οκτώβριος 2026')
            canvas.drawRightString(page_width - margin, 25, str(document.page))

    PDF.parent.mkdir(parents=True, exist_ok=True)
    document = SimpleDocTemplate(str(PDF), pagesize=A4, leftMargin=margin, rightMargin=margin, topMargin=57, bottomMargin=49, title='Property Co Pilot — Τεχνική τεκμηρίωση', author='')
    document.build(story, onFirstPage=decorate, onLaterPages=decorate)
    (WORK / 'expected-pages.json').write_text(json.dumps(expected_pages, ensure_ascii=False, indent=2), encoding='utf-8')
    print(f'PDF {PDF}')


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--format', choices=['docx', 'pdf', 'both'], default='both')
    arguments = parser.parse_args()
    if arguments.format in ('docx', 'both'):
        make_docx()
    if arguments.format in ('pdf', 'both'):
        make_pdf()
